from pathlib import Path
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Complete_EV_Charging_Analysis_Report.docx"
FIG = ROOT / "figures"

df = pd.read_csv(ROOT / "outputs/cleaned_ev_charging_data.csv", low_memory=False)
df["start_datetime"] = pd.to_datetime(df["start_datetime"], errors="coerce")
station_col = "Station Name"

total_sessions = len(df)
total_energy = df["energy_kWh"].sum()
mean_energy = df["energy_kWh"].mean()
median_energy = df["energy_kWh"].median()
mean_charge = df["charging_duration_hours"].mean()
median_charge = df["charging_duration_hours"].median()
mean_rate = df["energy_per_hour"].mean()
mean_idle = df["idle_time_hours"].mean()
median_idle = df["idle_time_hours"].median()
idle_pct = (df["idle_time_hours"] >= 1).mean() * 100
pearson_r = df[["charging_duration_hours", "energy_kWh"]].corr().iloc[0, 1]

day_order = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"]
hourly = df.groupby("hour_of_day").size()
dow = df.groupby("day_of_week").size().reindex(day_order)
month = df.groupby("month_name").size()
annual = df.groupby("year").agg(sessions=("energy_kWh","size"), energy_kWh=("energy_kWh","sum"), avg_energy=("energy_kWh","mean"), avg_hours=("charging_duration_hours","mean"))
ww = df.groupby("day_type").agg(sessions=("energy_kWh","size"), total_energy=("energy_kWh","sum"), avg_energy=("energy_kWh","mean"), median_energy=("energy_kWh","median"), avg_charge=("charging_duration_hours","mean"), median_charge=("charging_duration_hours","median"), avg_idle=("idle_time_hours","mean"))
season_order = ["Winter","Spring","Summer","Fall"]
seasonal = df.groupby("season").agg(sessions=("energy_kWh","size"), total_energy=("energy_kWh","sum"), avg_energy=("energy_kWh","mean"), avg_charge=("charging_duration_hours","mean")).reindex(season_order)
stations = pd.read_csv(ROOT / "outputs/station_summary.csv")
energy_q = df["energy_kWh"].quantile([.01,.25,.5,.75,.90,.95,.99])
duration_q = df["charging_duration_hours"].quantile([.01,.25,.5,.75,.90,.95,.99])

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.8)
styles["Normal"].font.color.rgb = RGBColor(0,0,0)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(24); styles["Title"].font.bold = True; styles["Title"].font.color.rgb = RGBColor(0,0,0)
for name, size in [("Heading 1",16),("Heading 2",13),("Heading 3",11.5)]:
    s = styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6); s.paragraph_format.keep_with_next=True

if "Figure Caption" not in styles:
    cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
else: cap = styles["Figure Caption"]
cap.font.name="Aptos"; cap.font.size=Pt(9); cap.font.italic=True; cap.font.color.rgb=RGBColor(55,55,55)
cap.paragraph_format.space_after=Pt(8); cap.paragraph_format.keep_with_next=False

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        el=tcMar.find(qn(f"w:{m}"))
        if el is None: el=OxmlElement(f"w:{m}"); tcMar.append(el)
        el.set(qn("w:w"),str(v)); el.set(qn("w:type"),"dxa")

def add_table(headers, rows, widths=None, font_size=9):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.style="Table Grid"; table.autofit=False
    for j,h in enumerate(headers):
        cell=table.rows[0].cells[j]; cell.text=str(h); set_cell_shading(cell,"1F4E78"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font_size)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"), "true"); trPr.append(tblHeader)
    for i,row in enumerate(rows):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            cells[j].text=str(val); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cells[j], "EAF2F8" if i%2 else "FFFFFF")
            for r in cells[j].paragraphs[0].runs: r.font.size=Pt(font_size); r.font.color.rgb=RGBColor(0,0,0)
    for row in table.rows:
        for j,cell in enumerate(row.cells):
            set_cell_margins(cell)
            if widths: cell.width=Inches(widths[j])
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_bullet(text):
    p=doc.add_paragraph(style="List Bullet"); p.add_run(text); p.paragraph_format.space_after=Pt(3)

def add_figure(name, caption, width=6.65):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    shape = p.add_run().add_picture(str(FIG/name), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("title", caption.split("  ", 1)[0])
    c=doc.add_paragraph(caption, style="Figure Caption"); c.alignment=WD_ALIGN_PARAGRAPH.CENTER

def page_break(): doc.add_page_break()

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run("Page ")
    fldChar1=OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"),"begin")
    instrText=OxmlElement("w:instrText"); instrText.set(qn("xml:space"),"preserve"); instrText.text="PAGE"
    fldChar2=OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"),"end")
    run._r.extend([fldChar1,instrText,fldChar2])

header=sec.header.paragraphs[0]; header.text="Exploratory Data Analysis of Electric Vehicle Charging Patterns"; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in header.runs: r.font.name="Aptos"; r.font.size=Pt(8); r.font.color.rgb=RGBColor(0,0,0)
add_page_number(sec.footer.paragraphs[0])
for r in sec.footer.paragraphs[0].runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(0,0,0)

# Title page
doc.add_paragraph().add_run("\n\n")
title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run("Exploratory Data Analysis of Electric Vehicle Charging Patterns")
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run("City of Palo Alto Electric Vehicle Charging Station Usage\nJuly 2011 to December 2020").bold=True
sub.runs[0].font.size=Pt(14)
doc.add_paragraph("\n")
meta=add_table(["Project Information","Details"],[
    ["Course context","University Data Science class project"],
    ["Analysis type","Exploratory data analysis"],
    ["Source file","EVChargingStationUsage.csv"],
    ["Original data","259,415 records and 33 columns"],
    ["Cleaned data","259,295 charging sessions"],
    ["Software","Python with pandas, NumPy, Matplotlib, Seaborn, SciPy, and scikit-learn"],
],[1.55,4.85],9.5)
doc.add_paragraph("Prepared as a reproducible analysis. The accompanying Jupyter notebook contains all code, executed outputs, figures, and exported summary datasets.").alignment=WD_ALIGN_PARAGRAPH.CENTER
page_break()

doc.add_heading("Abstract", level=1)
doc.add_paragraph(f"This project analyzes {total_sessions:,} cleaned electric vehicle charging sessions recorded at City of Palo Alto charging stations between July 29, 2011 and December 31, 2020. The analysis examines when chargers were used, how activity changed over time, how charging duration related to energy delivery, how stations differed, and whether weekday, weekend, seasonal, and idle-occupancy patterns were present. Cleaning removed 120 duplicate or objectively invalid records while retaining statistically extreme but internally valid sessions for investigation. The sessions delivered {total_energy:,.2f} kWh in total. A session delivered {mean_energy:.2f} kWh on average and involved {mean_charge:.2f} hours of active charging. Charging duration and energy had a Pearson correlation of {pearson_r:.3f}. Weekdays accounted for {(~df['weekend']).mean()*100:.1f}% of sessions, and {idle_pct:.1f}% of sessions remained connected for at least one hour after active charging ended. These findings describe observed behavior within one city and do not establish causal effects.")

doc.add_heading("Key Findings", level=1)
for x in [
    "Charging sessions most often started at 11 AM. Wednesday was the busiest day of the week, and January had the highest session count when calendar months were combined across years.",
    f"The cleaned records contained {total_energy:,.2f} kWh of delivered energy. Average energy was {mean_energy:.2f} kWh per session, and median energy was {median_energy:.2f} kWh.",
    f"Average active charging time was {mean_charge:.2f} hours, compared with a median of {median_charge:.2f} hours.",
    f"Charging duration and energy showed a strong positive linear association with r = {pearson_r:.3f}.",
    "The busiest year was 2017 with 48,950 sessions. The first year, 2011, contains only six months and is not directly comparable with full calendar years.",
    "PALO ALTO CA / HAMILTON #2 recorded the most sessions, with 23,706.",
    f"Average idle time was {mean_idle:.2f} hours, but the median was only {median_idle:.2f} hours. The difference indicates a right-skewed distribution with a smaller group of long idle events.",
]: add_bullet(x)

doc.add_heading("Contents", level=1)
for item in ["Introduction","Research Questions","Dataset and Variables","Data Cleaning and Feature Engineering","Descriptive Statistics","Results","Outlier and Correlation Analysis","Optional Clustering","Discussion","Limitations","Conclusion","References","Appendix"]:
    doc.add_paragraph(item, style="List Number")
page_break()

doc.add_heading("Introduction", level=1)
doc.add_paragraph("Public charging records provide a direct view of how often charging equipment is used, when sessions begin, how long vehicles actively charge, and how long they remain connected. These measures address different operational questions. Session counts describe recorded activity, energy measures electrical delivery, active charging duration describes equipment use for energy transfer, and total connection duration includes time when a vehicle may occupy a port without actively charging.")
doc.add_paragraph("This study uses session-level data published by the City of Palo Alto. Its purpose is to describe charging behavior and identify temporal, station-level, and occupancy patterns. The analysis does not attempt to predict future use or assign behavioral identities to drivers. It also does not treat correlations as causal evidence.")

doc.add_heading("Research Questions", level=1)
questions=[
    "When are EV chargers used most?",
    "How has charging usage changed over time?",
    "What is the relationship between charging duration and energy delivered?",
    "How do charging stations differ?",
    "How do weekday and weekend sessions compare?",
    "Do seasonal patterns appear in the data?",
    "How much time do vehicles remain connected after active charging ends?",
]
for q in questions: add_bullet(q)

doc.add_heading("Dataset and Variables", level=1)
doc.add_paragraph("The original CSV contains 259,415 rows and 33 columns. One row represents a recorded charging event or session. The source includes station name, start and end times, total connection duration, active charging time, delivered energy, greenhouse gas and gasoline savings estimates, port and plug characteristics, location fields, fee information, event identifiers, and equipment details.")
add_table(["Analytical concept","Source or engineered variable","Interpretation"],[
    ["Session start","Start Date","Beginning timestamp used for temporal features"],
    ["Session end","End Date","Ending timestamp used for validation"],
    ["Energy","Energy (kWh)","Electricity delivered during the session"],
    ["Active charging","Charging Time (hh:mm:ss)","Time during which charging was active"],
    ["Connection time","Total Duration (hh:mm:ss)","Total time the vehicle occupied the port"],
    ["Idle time","Connection time minus active charging","Estimated post or interrupted charging occupancy"],
    ["Charging rate","Energy divided by active charging time","Average kW equivalent across the session"],
    ["Station","Station Name","Named charging location or unit"],
],[1.25,2.2,3.0],8.6)
page_break()

doc.add_heading("Data Cleaning and Feature Engineering", level=1)
doc.add_heading("Cleaning Rules", level=2)
doc.add_paragraph("The cleaning process removed only records that could not support a valid session-level analysis. Statistical outliers were not removed solely because they were extreme.")
add_table(["Check","Action"],[
    ["Exact duplicate row","Keep the first occurrence and remove later duplicates"],
    ["Invalid start or end timestamp","Remove because the session cannot be placed in time"],
    ["End earlier than start","Remove as temporally inconsistent"],
    ["Negative or missing energy","Remove because delivered energy is invalid or unavailable"],
    ["Zero or negative duration","Remove because rate and duration analyses would be invalid"],
    ["Charging time materially longer than connection time","Remove as internally inconsistent"],
    ["Missing descriptive field","Retain when core session variables remain valid"],
    ["Extreme valid measurement","Retain and examine in the outlier section"],
],[2.45,4.0],9)
doc.add_paragraph("The procedure removed 120 records, leaving 259,295 sessions. Small negative idle differences of at most one second, consistent with timestamp rounding, were set to zero after the validity checks.")

doc.add_heading("Engineered Features", level=2)
doc.add_paragraph("The start timestamp generated year, numeric month, month name, day of week, hour of day, weekend status, season, calendar date, and year-month. Durations were converted from hour-minute-second strings to decimal hours. Idle time was calculated as connection duration minus active charging duration. Energy per charging hour was calculated only when active charging duration was positive. Cost per kWh was calculated where fee and positive energy were present.")

doc.add_heading("Descriptive Statistics", level=1)
add_table(["Measure","Result"],[
    ["Cleaned sessions",f"{total_sessions:,}"],
    ["Total energy",f"{total_energy:,.2f} kWh"],
    ["Average energy per session",f"{mean_energy:.2f} kWh"],
    ["Median energy per session",f"{median_energy:.2f} kWh"],
    ["Average active charging duration",f"{mean_charge:.2f} hours"],
    ["Median active charging duration",f"{median_charge:.2f} hours"],
    ["Average energy per active charging hour",f"{mean_rate:.2f} kW equivalent"],
    ["Unique stations",f"{df[station_col].nunique():,}"],
    ["Date range","July 29, 2011 to December 31, 2020"],
],[3.6,2.0],9.4)
doc.add_heading("Selected Percentiles", level=2)
rows=[]
for q in [.01,.25,.5,.75,.90,.95,.99]: rows.append([f"{q*100:.0f}th",f"{energy_q[q]:.2f}",f"{duration_q[q]:.2f}"])
add_table(["Percentile","Energy kWh","Charging duration hours"],rows,[1.4,2.1,2.3],9.2)
page_break()

doc.add_heading("Results", level=1)
doc.add_heading("When Chargers Were Used Most", level=2)
doc.add_paragraph("Session starts concentrated in daytime hours. The highest count occurred at 11 AM. Wednesday had the largest day-of-week count, January had the largest combined calendar-month count, and weekdays represented 76.1 percent of sessions. These counts show recorded use. They do not measure queues, failed charging attempts, or unmet demand.")
add_figure("charging_sessions_by_time.png","Figure 1  Charging session counts by hour, day of week, calendar month, and day type",6.55)

doc.add_heading("Changes in Usage Over Time", level=2)
doc.add_paragraph("Monthly session count and total energy increased substantially during the earlier years, although the series also contains declines and recoveries. The busiest full year was 2017 with 48,950 sessions. May 2017 was the busiest individual month with 4,982 sessions. Average energy per session and average charging duration changed less smoothly than volume. Calendar year 2011 contains only July through December, so its annual total is incomplete.")
add_figure("monthly_charging_trend.png","Figure 2  Monthly sessions, total energy, average energy, and average charging duration",6.55)
page_break()

doc.add_heading("Charging Duration and Energy", level=2)
doc.add_paragraph(f"Active charging duration and energy delivered had a Pearson correlation of {pearson_r:.3f}. The p value was effectively below conventional reporting precision because of the large sample. The positive association is clear, but the scatter also shows variation at similar durations. Extremely long sessions did not always deliver proportionally high energy. Vehicle charging limits, charger power, battery state, interruptions, and data-recording behavior may contribute to this spread.")
add_figure("duration_vs_energy.png","Figure 3  Charging duration and energy delivered with a fitted linear trend",6.55)

doc.add_heading("Differences Among Stations", level=2)
doc.add_paragraph("Stations differed in session count, total energy, average energy, charging duration, and idle time. PALO ALTO CA / HAMILTON #2 recorded 23,706 sessions, the most in the dataset. The ordering by total energy differed from the ordering by sessions, which indicates that session volume alone does not fully describe station use.")
add_figure("top_charging_stations.png","Figure 4  Leading stations by session count and total delivered energy",6.55)
top_rows=[]
for _,r in stations.head(10).iterrows(): top_rows.append([r[station_col],f"{int(r.sessions):,}",f"{r.total_energy_kWh:,.0f}",f"{r.average_energy_kWh:.2f}",f"{r.average_charging_hours:.2f}"])
add_table(["Station","Sessions","Total kWh","Mean kWh","Mean hours"],top_rows,[2.55,0.78,0.88,0.83,0.86],7.8)
page_break()

doc.add_heading("Weekday and Weekend Charging", level=2)
weekday=ww.loc["Weekday"]; weekend=ww.loc["Weekend"]
energy_diff=(weekend.avg_energy/weekday.avg_energy-1)*100; duration_diff=(weekend.avg_charge/weekday.avg_charge-1)*100
doc.add_paragraph(f"Weekday sessions accounted for 76.1 percent of the dataset. Weekend sessions delivered {abs(energy_diff):.1f} percent less average energy and had {abs(duration_diff):.1f} percent shorter average active charging time than weekday sessions. Start-time profiles also differed. These comparisons are unadjusted for year, station composition, or charger type.")
add_table(["Day type","Sessions","Total energy kWh","Mean energy kWh","Mean charge hours","Mean idle hours"],[
    ["Weekday",f"{int(weekday.sessions):,}",f"{weekday.total_energy:,.0f}",f"{weekday.avg_energy:.2f}",f"{weekday.avg_charge:.2f}",f"{weekday.avg_idle:.2f}"],
    ["Weekend",f"{int(weekend.sessions):,}",f"{weekend.total_energy:,.0f}",f"{weekend.avg_energy:.2f}",f"{weekend.avg_charge:.2f}",f"{weekend.avg_idle:.2f}"],
],[1.0,0.82,1.2,1.08,1.18,1.0],8.2)
add_figure("weekday_weekend_comparison.png","Figure 5  Weekday and weekend energy, duration, and start-time patterns",6.55)

doc.add_heading("Seasonal Patterns", level=2)
doc.add_paragraph("Winter recorded the largest number of sessions and the highest average energy per session. Seasonal totals combine observations from multiple years and reflect the changing number and availability of stations. The evidence therefore does not isolate a weather effect.")
season_rows=[]
for s,r in seasonal.iterrows(): season_rows.append([s,f"{int(r.sessions):,}",f"{r.total_energy:,.0f}",f"{r.avg_energy:.2f}",f"{r.avg_charge:.2f}"])
add_table(["Season","Sessions","Total energy kWh","Mean energy kWh","Mean charge hours"],season_rows,[1.0,1.0,1.35,1.25,1.25],8.7)
add_figure("seasonal_charging_patterns.png","Figure 6  Session count and average energy per session by season",6.55)
page_break()

doc.add_heading("Occupancy and Idle Behavior", level=2)
doc.add_paragraph(f"Idle time measured the difference between total connection time and active charging time. Its mean was {mean_idle:.2f} hours, while its median was {median_idle:.2f} hours. The large difference between these measures reflects a right-skewed distribution. Using one hour as an analytical threshold, {idle_pct:.1f} percent of sessions had substantial idle time. This threshold is not an official standard and does not indicate whether a vehicle violated a parking rule.")
add_figure("idle_time_distribution.png","Figure 7  Idle-time distribution through the 99th percentile",6.55)
add_figure("stations_highest_idle_time.png","Figure 8  Stations with the highest average idle time among stations with at least 100 sessions",6.4)

doc.add_heading("Outlier and Correlation Analysis", level=1)
doc.add_heading("Potential Outliers", level=2)
doc.add_paragraph("The analysis used interquartile-range fences and 99th percentiles to identify unusual energy, duration, connection, idle, and rate values. These rules flag records for review but do not prove that the records are errors. Long connections may reflect overnight parking, interrupted charging, post-charge occupancy, or equipment behavior. High energy or rate values may reflect charger and vehicle differences.")
add_figure("outlier_boxplots.png","Figure 9  Box plots for energy, active charging duration, and energy per active charging hour",6.55)
page_break()

doc.add_heading("Correlations", level=2)
doc.add_paragraph("The correlation matrix included energy, active charging duration, connection duration, idle duration, fee, energy per hour, and start hour where meaningful. Numeric identifiers such as event ID, station equipment ID, and port number were excluded because arithmetic differences between identifier values have no analytical interpretation.")
add_figure("correlation_heatmap.png","Figure 10  Correlations among meaningful numerical variables",6.25)
doc.add_paragraph("The strongest substantive relationship was between active charging duration and energy. Connection duration also related strongly to active charging duration because connection time contains both active and idle periods. These correlations summarize linear association and do not identify causes.")

doc.add_heading("Optional Clustering", level=1)
doc.add_paragraph("A small exploratory K-means analysis examined active charging duration, energy, idle time, and energy per active charging hour. The variables were standardized so that their different units did not determine the clusters. Values were winsorized at the first and 99th percentiles for clustering stability, but the cleaned analytical dataset was not altered. Candidate solutions from two through five clusters were compared using inertia and silhouette score.")
add_figure("optional_clustering_diagnostics.png","Figure 11  Elbow and silhouette diagnostics for candidate K-means solutions",6.55)
doc.add_paragraph("The selected solution was interpreted through its measured cluster statistics. No labels such as commuter, resident, or heavy user were assigned because the dataset does not establish trip purpose, residence, or user intent.")
page_break()

doc.add_heading("Discussion", level=1)
doc.add_paragraph("The analysis shows that charging demand is structured by time and location. Daytime session starts and the weekday majority are consistent with a network that serves substantial weekday activity, although the data do not identify the purposes of those trips. Station rankings also show concentration: a limited number of stations handled many sessions, while energy rankings varied because sessions differed in delivered energy.")
doc.add_paragraph("Active charging time and connection time should not be treated as interchangeable. Most sessions had little idle time, as indicated by the near-zero median, but the average was pulled upward by longer idle events. This combination matters operationally because a modest subset of sessions can occupy ports well after energy delivery slows or ends.")
doc.add_paragraph("The strong duration-energy relationship confirms that active charging time is useful for explaining energy delivery. The remaining spread is equally important. Sessions with similar duration can deliver different energy, so analyses of electricity demand should retain energy as a direct measure rather than substituting duration alone.")
doc.add_paragraph("Temporal comparisons require attention to changing exposure. The dataset spans more than nine years, during which station inventory and charging behavior may have changed. Annual and seasonal totals combine periods with different infrastructure and activity levels. The results therefore support descriptive comparisons but not causal claims about why use rose or fell.")

doc.add_heading("Limitations", level=1)
for x in [
    "The records describe one geographic area and may not generalize to other cities or charging networks.",
    "The dataset is observational. Correlation and temporal coincidence do not establish causation.",
    "Station inventory, charger type, availability, pricing, and parking policies may have changed during the study period.",
    "Calendar year 2011 contains only six months, which makes its annual total incomplete.",
    "Several descriptive equipment and user fields contain missing values.",
    "Statistically extreme records may represent valid behavior or undetected logging errors.",
    "The data do not directly measure queues, failed attempts, unmet demand, state of charge, trip purpose, or driver intent.",
    "A station name may represent a location, unit, or port configuration rather than a perfectly comparable facility.",
]: add_bullet(x)

doc.add_heading("Conclusion", level=1)
doc.add_paragraph(f"The Palo Alto records provide a detailed view of {total_sessions:,} valid charging sessions across 47 stations. Activity concentrated in daytime weekday periods, and station use was uneven. Sessions delivered {total_energy:,.0f} kWh in total, with mean energy of {mean_energy:.2f} kWh. Active charging duration and energy had a strong positive association, but duration did not fully determine energy delivery. Idle behavior was usually small, yet {idle_pct:.1f} percent of sessions remained connected for at least one hour beyond active charging. These findings show why charging systems should be described with session count, energy, active charging time, connection time, and station-level measures together.")

doc.add_heading("References", level=1)
refs=[
    "City of Palo Alto. Electric Vehicle Charging Station Usage July 2011 to December 2020. City of Palo Alto Open Data Portal. https://data.paloalto.gov/datasets/194693/electric-vehicle-charging-station-usage-july-2011-dec-2020/",
    "Harris, C. R., et al. 2020. Array programming with NumPy. Nature 585, 357 to 362.",
    "Hunter, J. D. 2007. Matplotlib A 2D graphics environment. Computing in Science and Engineering 9, 90 to 95.",
    "McKinney, W. 2010. Data structures for statistical computing in Python. Proceedings of the 9th Python in Science Conference, 56 to 61.",
    "Pedregosa, F., et al. 2011. Scikit-learn Machine learning in Python. Journal of Machine Learning Research 12, 2825 to 2830.",
    "Virtanen, P., et al. 2020. SciPy 1.0 Fundamental algorithms for scientific computing in Python. Nature Methods 17, 261 to 272.",
    "Waskom, M. L. 2021. Seaborn Statistical data visualization. Journal of Open Source Software 6, 3021.",
]
for ref in refs:
    p=doc.add_paragraph(ref); p.paragraph_format.left_indent=Inches(.3); p.paragraph_format.first_line_indent=Inches(-.3); p.paragraph_format.space_after=Pt(6)

page_break()
doc.add_heading("Appendix", level=1)
doc.add_heading("Generated Project Files", level=2)
for x in [
    "EV_Charging_Analysis.ipynb contains the full executed analysis.",
    "PROJECT_RESULTS.md contains the concise results summary.",
    "EV_Charging_Project_Report.pdf contains the earlier visual report.",
    "outputs contains the cleaned data and hourly, daily, monthly, station, and outlier summaries.",
    "figures contains all high-resolution charts used in the analysis and reports.",
]: add_bullet(x)
doc.add_heading("Reproducibility Notes", level=2)
doc.add_paragraph("The notebook automatically selects the EV charging CSV from its project folder, identifies relevant source columns by normalized names, creates the output directories, performs cleaning and feature engineering, produces summary files and figures, and writes the Markdown results summary. All notebook code cells were executed successfully before this report was prepared.")

doc.core_properties.title="Exploratory Data Analysis of Electric Vehicle Charging Patterns"
doc.core_properties.subject="Complete university data science project report"
doc.core_properties.author="University Data Science Project"
doc.save(OUT)
print(OUT)
